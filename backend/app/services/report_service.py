"""Report generation: PDF (ReportLab), CSV and XLSX.

Layout per spec — header with logo placeholder + company, executive summary,
colour-coded tender listing, "Powered by TenderRadar" footer.
"""
from __future__ import annotations

import csv
import logging
import os
from datetime import date, datetime, timedelta

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import PortalSource, Report, Tender, TenderKeywordMatch, User
from ..utils.timeutil import fmt_ist, ist_now

logger = logging.getLogger(__name__)

NAVY = colors.HexColor("#0F1629")
TEAL = colors.HexColor("#0A9396")
ROW_HIGH = colors.HexColor("#d9f2e5")
ROW_MED = colors.HexColor("#fdf3d7")
ROW_LOW = colors.HexColor("#fbe3df")


def _report_dir() -> str:
    path = os.path.join(settings.storage_dir, "reports")
    os.makedirs(path, exist_ok=True)
    return path


def _query_tenders(db: Session, user: User, date_from: date, date_to: date,
                   filters: dict) -> list[dict]:
    """User's matched tenders within range, best score per tender."""
    stmt = (
        select(Tender, func.max(TenderKeywordMatch.relevance_score).label("score"),
               PortalSource.name, PortalSource.code)
        .join(TenderKeywordMatch, TenderKeywordMatch.tender_id == Tender.id)
        .join(PortalSource, Tender.portal_source_id == PortalSource.id)
        .where(
            TenderKeywordMatch.user_id == user.id,
            Tender.is_archived.is_(False),
            func.date(Tender.scraped_at) >= date_from,
            func.date(Tender.scraped_at) <= date_to,
        )
        .group_by(Tender.id, PortalSource.name, PortalSource.code)
    )
    if filters.get("portal_code"):
        stmt = stmt.where(PortalSource.code == filters["portal_code"])
    if filters.get("category"):
        stmt = stmt.where(Tender.category == filters["category"])
    if filters.get("state"):
        stmt = stmt.where(Tender.state == filters["state"])
    if filters.get("min_score") is not None:
        stmt = stmt.having(func.max(TenderKeywordMatch.relevance_score) >= filters["min_score"])
    stmt = stmt.order_by(func.max(TenderKeywordMatch.relevance_score).desc())

    rows = db.execute(stmt).all()
    return [
        {
            "ref": t.tender_ref_no, "title": t.title,
            "organisation": t.organisation or "—",
            "category": (t.category or "—").title(),
            "published": t.published_date.strftime("%d %b %Y") if t.published_date else "—",
            "closing": fmt_ist(t.closing_date, "%d %b %Y %I:%M %p"),
            "value": _fmt_inr(t.estimated_value),
            "portal": portal_name, "portal_code": portal_code,
            "score": round(float(score or 0), 1),
            "state": t.state or "—",
        }
        for t, score, portal_name, portal_code in rows
    ]


def _fmt_inr(value) -> str:
    if value is None:
        return "—"
    value = float(value)
    if value >= 1e7:
        return f"₹{value / 1e7:.2f} Cr"
    if value >= 1e5:
        return f"₹{value / 1e5:.2f} L"
    return f"₹{value:,.0f}"


def generate_report(db: Session, user: User, *, report_type: str,
                    file_format: str = "pdf", date_from: date | None = None,
                    date_to: date | None = None, filters: dict | None = None) -> Report:
    filters = filters or {}
    today = ist_now().date()
    if report_type == "daily":
        date_from = date_to = today
    elif report_type == "weekly":
        date_to = today
        date_from = today - timedelta(days=7)
    else:  # custom
        date_from = date_from or today - timedelta(days=30)
        date_to = date_to or today

    tenders = _query_tenders(db, user, date_from, date_to, filters)
    ts = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    filename = f"tenderradar_{report_type}_{user.id}_{ts}.{file_format}"
    file_path = os.path.join(_report_dir(), filename)

    if file_format == "pdf":
        _build_pdf(file_path, user, report_type, date_from, date_to, tenders)
    elif file_format == "csv":
        _build_csv(file_path, tenders)
    elif file_format == "xlsx":
        _build_xlsx(file_path, tenders)
    else:
        raise ValueError(f"Unsupported format: {file_format}")

    report = Report(
        user_id=user.id, report_type=report_type, file_format=file_format,
        date_range_start=date_from, date_range_end=date_to,
        filters_json=filters, file_path=file_path, tender_count=len(tenders),
    )
    db.add(report)
    db.commit()
    db.refresh(report)
    return report


# --- PDF ---------------------------------------------------------------------

def _footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 10 * mm,
                      "Powered by TenderRadar · This report aggregates publicly available tender "
                      "information. Always verify on the official portal before bidding.")
    canvas.drawRightString(landscape(A4)[0] - 15 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def _build_pdf(path: str, user: User, report_type: str,
               date_from: date, date_to: date, tenders: list[dict]) -> None:
    doc = SimpleDocTemplate(
        path, pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=18 * mm,
        title=f"TenderRadar {report_type.title()} Report",
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], textColor=NAVY,
                        fontSize=18, alignment=0, spaceAfter=2)
    sub = ParagraphStyle("sub", parent=styles["Normal"], textColor=colors.grey, fontSize=9)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=TEAL, fontSize=12)
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7.5, leading=9)

    titles = {"daily": "Daily Summary Report", "weekly": "Weekly Intelligence Report",
              "custom": "Custom Tender Report"}
    story = []

    # Header — [logo placeholder] company + meta
    logo = Table([["📡 TenderRadar"]], colWidths=[45 * mm], rowHeights=[12 * mm])
    logo.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROUNDEDCORNERS", (0, 0), (-1, -1), [4, 4, 4, 4]),
    ]))
    header = Table(
        [[logo,
          Table([[Paragraph(titles.get(report_type, "Tender Report"), h1)],
                 [Paragraph(f"{user.company_name or user.email} &nbsp;·&nbsp; "
                            f"{date_from:%d %b %Y} – {date_to:%d %b %Y} &nbsp;·&nbsp; "
                            f"Generated {ist_now():%d %b %Y, %I:%M %p} IST", sub)]],
                style=[("LEFTPADDING", (0, 0), (-1, -1), 0)])]],
        colWidths=[50 * mm, None],
    )
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                ("LEFTPADDING", (0, 0), (0, 0), 0)]))
    story += [header, Spacer(1, 6 * mm)]

    # Executive summary
    by_portal: dict[str, int] = {}
    by_category: dict[str, int] = {}
    high = sum(1 for t in tenders if t["score"] >= 70)
    for t in tenders:
        by_portal[t["portal"]] = by_portal.get(t["portal"], 0) + 1
        by_category[t["category"]] = by_category.get(t["category"], 0) + 1
    top_portals = ", ".join(f"{k} ({v})" for k, v in
                            sorted(by_portal.items(), key=lambda x: -x[1])[:4]) or "—"
    cats = ", ".join(f"{k} ({v})" for k, v in
                     sorted(by_category.items(), key=lambda x: -x[1])) or "—"

    story.append(Paragraph("Executive Summary", h2))
    summary = Table([
        ["Total tenders found", str(len(tenders)),
         "High relevance (70+)", str(high)],
        ["Most active portals", top_portals, "By category", cats],
    ], colWidths=[42 * mm, 50 * mm, 42 * mm, None])
    summary.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef1f7")),
        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#eef1f7")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#c9d0de")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story += [summary, Spacer(1, 6 * mm), Paragraph("Tender Listing", h2)]

    # Listing table (colour-coded by relevance band)
    head = ["Ref No", "Title", "Organisation", "Category", "Published",
            "Closing Date", "Est. Value", "Portal", "Score"]
    data = [head]
    for t in tenders:
        data.append([
            Paragraph(t["ref"], cell), Paragraph(t["title"][:220], cell),
            Paragraph(t["organisation"][:120], cell), t["category"],
            t["published"], t["closing"], t["value"],
            Paragraph(t["portal"][:60], cell), f'{t["score"]:.0f}',
        ])
    if len(data) == 1:
        data.append(["—", "No tenders matched this period", "", "", "", "", "", "", ""])

    table = Table(data, repeatRows=1,
                  colWidths=[34 * mm, 70 * mm, 42 * mm, 17 * mm, 19 * mm,
                             27 * mm, 18 * mm, 28 * mm, 12 * mm])
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c9d0de")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i, t in enumerate(tenders, start=1):
        band = ROW_HIGH if t["score"] >= 70 else ROW_MED if t["score"] >= 40 else ROW_LOW
        style.append(("BACKGROUND", (0, i), (-1, i), band))
    table.setStyle(TableStyle(style))
    story.append(table)

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)


# --- CSV / XLSX -----------------------------------------------------------------

_COLUMNS = ["ref", "title", "organisation", "category", "state", "published",
            "closing", "value", "portal", "score"]
_HEADERS = ["Ref No", "Title", "Organisation", "Category", "State", "Published",
            "Closing Date", "Est. Value", "Portal", "Relevance Score"]


def _build_csv(path: str, tenders: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(_HEADERS)
        for t in tenders:
            writer.writerow([t[c] for c in _COLUMNS])


def _build_xlsx(path: str, tenders: list[dict]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Tenders"
    ws.append(_HEADERS)
    header_fill = PatternFill("solid", fgColor="0F1629")
    for c in ws[1]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = header_fill
    fills = {"high": PatternFill("solid", fgColor="D9F2E5"),
             "med": PatternFill("solid", fgColor="FDF3D7"),
             "low": PatternFill("solid", fgColor="FBE3DF")}
    for t in tenders:
        ws.append([t[c] for c in _COLUMNS])
        band = "high" if t["score"] >= 70 else "med" if t["score"] >= 40 else "low"
        for c in ws[ws.max_row]:
            c.fill = fills[band]
    widths = [22, 60, 35, 12, 16, 14, 22, 14, 28, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    wb.save(path)
